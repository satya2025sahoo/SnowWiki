import os
import re
import uuid
import logging
import time
from docx import Document

from src.ingestion.docx_utils.heuristics import _is_valid_pseudo_heading, _is_servicenow_code
from src.ingestion.docx_utils.admonitions import _ADMONITION_RE, _ADMONITION_MAP, _format_admonition
from src.ingestion.docx_utils.tables import _convert_table_to_markdown_or_html
from src.ingestion.docx_utils.extraction import _extract_paragraph_markdown
from src.ingestion.docx_utils.normalization import _normalize_docx_markdown

logger = logging.getLogger(__name__)


def _fallback_convert_docx_to_markdown(file_path: str, branch_name: str = "default") -> str:
    """
    Pure-python DOCX → Markdown conversion used when mammoth is unavailable.
    Preserves hyperlinks, detects admonitions, serializes complex tables,
    and extracts embedded images.
    """
    doc = Document(file_path)
    md_lines = []

    topic_pattern = re.compile(r"^(Topic\s+\d+[:\-]?.*)$", re.IGNORECASE)
    numbered_pattern = re.compile(r"^(\d+\.\s+[A-Z][a-zA-Z0-9\s&]+)$")

    prefix = os.path.splitext(os.path.basename(file_path))[0]
    image_dir = None
    try:
        from src.ingestion.core import get_branch_image_dir
        image_dir = get_branch_image_dir(branch_name)
    except Exception:
        pass

    image_describer = None
    try:
        from src.ingestion.image_describer import ImageDescriberAgent
        image_describer = ImageDescriberAgent()
    except Exception:
        pass

    for p in doc.paragraphs:
        style_name = p.style.name if p.style else ""
        text, is_bold_run = _extract_paragraph_markdown(
            p, branch_name=branch_name, image_dir=image_dir, image_describer=image_describer, prefix=prefix
        )
        if not text:
            continue

        raw_plain = p.text.strip()

        # Check callout / admonition first
        adm_match = _ADMONITION_RE.match(raw_plain)
        if adm_match:
            tag = _ADMONITION_MAP.get(adm_match.group(1).lower(), "NOTE")
            rest = adm_match.group(2).strip()
            md_lines.append("\n".join(_format_admonition(tag, rest)) + "\n")
            continue

        if style_name.startswith("Heading"):
            try:
                level = int(style_name.split()[-1])
                md_lines.append(f"{'#' * level} {raw_plain}\n")
            except ValueError:
                md_lines.append(f"# {raw_plain}\n")
        elif style_name == "Title":
            md_lines.append(f"# {raw_plain}\n")
        elif topic_pattern.match(raw_plain) or numbered_pattern.match(raw_plain):
            md_lines.append(f"## {raw_plain}\n")
        elif is_bold_run and _is_valid_pseudo_heading(raw_plain):
            md_lines.append(f"### {raw_plain}\n")
        elif style_name.startswith("List Bullet") or style_name == "List Paragraph":
            md_lines.append(f"- {text}")
        elif style_name.startswith("List Number"):
            md_lines.append(f"1. {text}")
        else:
            if _is_servicenow_code(raw_plain):
                md_lines.append(f"```javascript\n{raw_plain}\n```\n")
            else:
                md_lines.append(text + "\n")

    for table in doc.tables:
        table_md = _convert_table_to_markdown_or_html(table)
        if table_md:
            md_lines.append(table_md)

    raw_output = "\n".join(md_lines)
    return _normalize_docx_markdown(raw_output)


def convert_docx_to_markdown(file_path: str, branch_name: str = "default") -> str:
    """
    Convert a DOCX file to structured Markdown using Mammoth.
    Extracts embedded images, fetches LLM descriptions with robust fallbacks,
    preserves semantic hierarchy, and formats callouts and tables.
    Falls back to a pure-python implementation if mammoth is not installed.
    """
    try:
        import mammoth
        from src.ingestion.image_describer import ImageDescriberAgent
        from src.ingestion.core import get_branch_image_dir

        image_dir = get_branch_image_dir(branch_name)
        prefix = os.path.splitext(os.path.basename(file_path))[0]

        class ImageCounter:
            def __init__(self):
                self.val = 0

        counter = ImageCounter()
        image_describer = ImageDescriberAgent()

        def _sanitize_alt(desc: str) -> str:
            desc = desc.replace("\n", " ").replace("\r", " ")
            desc = desc.replace("[", "(").replace("]", ")")
            desc = re.sub(r"\s+", " ", desc).strip()
            return desc[:300]

        def convert_image(image):
            nonlocal counter
            i = counter.val
            counter.val += 1

            ext = ".png"
            if image.content_type == "image/jpeg":
                ext = ".jpg"
            elif image.content_type == "image/gif":
                ext = ".gif"

            filename = f"{prefix}_inline_{uuid.uuid4().hex[:6]}{ext}"
            out_path = os.path.join(image_dir, filename)
            rel_path = f"images/{filename}"

            with image.open() as image_bytes:
                with open(out_path, "wb") as f:
                    f.write(image_bytes.read())

            desc = None
            try:
                if image_describer:
                    if i > 0:
                        time.sleep(1)  # Light throttle to prevent 429 rate limit spikes
                    desc = image_describer.describe_from_path(out_path)
            except Exception as e:
                logger.warning(f"Error generating description for image {filename}: {e}")

            # Refinement 3: Image Alt-Text Fallback
            if not desc or "unavailable" in desc.lower() or "error" in desc.lower() or not desc.strip():
                desc = f"Image: {filename} - Description Unavailable"

            return {"src": rel_path, "alt": f"Image - Description: {_sanitize_alt(desc)}"}

        # Convert to Markdown using Mammoth
        style_map = (
            "p[style-name='Heading 1'] => h1:fresh\n"
            "p[style-name='Heading 2'] => h2:fresh\n"
            "p[style-name='Heading 3'] => h3:fresh\n"
            "p[style-name='Title'] => h1:fresh\n"
            "p[style-name='Subtitle'] => h2:fresh\n"
            "p[style-name='Code'] => pre:fresh\n"
            "p[style-name='HTML Preformatted'] => pre:fresh\n"
            "p[style-name='Source Code'] => pre:fresh"
        )
        with open(file_path, "rb") as docx_file:
            result = mammoth.convert_to_markdown(
                docx_file,
                style_map=style_map,
                convert_image=mammoth.images.img_element(convert_image)
            )
            raw_md = result.value

        normalized_md = _normalize_docx_markdown(raw_md)
        return normalized_md

    except ImportError:
        logger.warning("mammoth package not found. Using fallback DOCX to Markdown converter.")
        return _fallback_convert_docx_to_markdown(file_path, branch_name=branch_name)
